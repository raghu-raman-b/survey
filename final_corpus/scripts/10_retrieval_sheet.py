#!/usr/bin/env python3
"""Build the full-text retrieval tracker from the included studies.

Input:  study_corpus/study_corpus.json   (written by 09_combine.py)
Output: retrieval/retrieval_tracker.xlsx
        retrieval/retrieval_split.json   which corpus IDs went to whom

The included records are split into two random halves, stratified by publication
year, with a fixed seed so the split is reproducible. Each half gets its own sheet
with a status dropdown; the Summary sheet counts the statuses with live formulas,
giving the PRISMA "Reports sought for retrieval" / "Reports not retrieved" numbers.

Rerunning overwrites the workbook, so do not rerun once retrieval has started.

Usage:  python3 10_retrieval_sheet.py [--force]
"""
import json, os, random, sys, collections, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

OUT   = '/home/prex_san/Documents/survey/final_corpus'
SRC   = os.path.join(OUT, 'study_corpus', 'study_corpus.json')
RDIR  = os.path.join(OUT, 'retrieval')
XLSX  = os.path.join(RDIR, 'retrieval_tracker.xlsx')
SEED  = 20260928
PEOPLE = ['Nikhil', 'Raghu']
STATUSES = ['Not acquired', 'Acquired', 'Not retrievable']

if os.path.exists(XLSX) and '--force' not in sys.argv:
    sys.exit(f"{XLSX} already exists and may hold retrieval progress.\n"
             f"Pass --force to overwrite it.")

recs = json.load(open(SRC))
cid = lambda r: f"C{r['cid']:04d}"

# stratified random split: shuffle within each year, deal alternately
rng = random.Random(SEED)
by_year = collections.defaultdict(list)
for r in sorted(recs, key=cid):
    by_year[r['year']].append(r)
halves = {p: [] for p in PEOPLE}
turn = 0
for y in sorted(by_year):
    group = by_year[y]
    rng.shuffle(group)
    for r in group:
        halves[PEOPLE[turn % 2]].append(r); turn += 1

def link(r):
    doi = (r.get('doi_raw') or r.get('doi') or '').strip()
    if doi:
        return f"https://doi.org/{doi}", doi
    q = urllib.parse.quote_plus(r['title'])
    return f"https://scholar.google.com/scholar?q={q}", "(no DOI) search title"

HEAD = Font(bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='2F5597')
LINK = Font(color='0563C1', underline='single')

wb = Workbook()
summ = wb.active; summ.title = 'Summary'

for p in PEOPLE:
    ws = wb.create_sheet(p)
    cols = [('corpus_id', 10), ('title', 70), ('authors', 28), ('year', 6), ('venue', 34),
            ('doi / link', 28), ('status', 15), ('zotero', 8), ('notes', 40)]
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(1, i, h); c.font = HEAD; c.fill = HFILL
        ws.column_dimensions[c.column_letter].width = w
    rows = sorted(halves[p], key=lambda r: (r['year'], cid(r)))
    for i, r in enumerate(rows, 2):
        url, text = link(r)
        authors = '; '.join(r.get('authors', [])[:3]) + (' et al.' if len(r.get('authors', [])) > 3 else '')
        ws.append([cid(r), r['title'], authors, int(r['year']), r.get('venue', ''),
                   text, 'Not acquired', '', ''])
        ws.cell(i, 6).hyperlink = url; ws.cell(i, 6).font = LINK
        ws.cell(i, 2).alignment = Alignment(wrap_text=False)
    last = len(rows) + 1
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(f"G2:G{last}")
    dz = DataValidation(type='list', formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(dz); dz.add(f"H2:H{last}")
    green = PatternFill('solid', fgColor='C6EFCE'); red = PatternFill('solid', fgColor='FFC7CE')
    ws.conditional_formatting.add(f"G2:G{last}", CellIsRule(operator='equal', formula=['"Acquired"'], fill=green))
    ws.conditional_formatting.add(f"G2:G{last}", CellIsRule(operator='equal', formula=['"Not retrievable"'], fill=red))
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f"A1:I{last}"

# Summary sheet with live counts
summ.column_dimensions['A'].width = 34
for i in range(len(PEOPLE) + 1):
    summ.column_dimensions['BCD'[i]].width = 14
summ.append(['Full-text retrieval'] ); summ['A1'].font = Font(bold=True, size=14)
summ.append([])
summ.append(['Status'] + PEOPLE + ['Total'])
for c in summ[3]: c.font = HEAD; c.fill = HFILL
for s in STATUSES:
    row = [s] + [f'=COUNTIF(\'{p}\'!G:G,"{s}")' for p in PEOPLE]
    summ.append(row)
    r = summ.max_row
    summ.cell(r, len(PEOPLE) + 2, f"=SUM(B{r}:{'BCD'[len(PEOPLE)-1]}{r})")
summ.append(['Assigned'] + [len(halves[p]) for p in PEOPLE] + [f"=SUM(B{summ.max_row+1}:C{summ.max_row+1})"])
summ.cell(summ.max_row, 1).font = Font(bold=True)
summ.append([])
summ.append(['For the PRISMA flow']); summ.cell(summ.max_row, 1).font = Font(bold=True)
summ.append(['Reports sought for retrieval', len(recs)])
summ.append(['Reports not retrieved', "=D6"])
summ.append(['Still to look for', "=D4"])
summ.append(['Reports assessed for eligibility', "=D5"])
summ.append([])
summ.append(['Statuses: "Not acquired" = not found yet; "Acquired" = full text saved in Zotero;'])
summ.append(['"Not retrievable" = searched and could not get it (counts as PRISMA "not retrieved").'])
summ.append([f'Split: random, stratified by year, seed {SEED} (scripts/10_retrieval_sheet.py).'])

os.makedirs(RDIR, exist_ok=True)
wb.save(XLSX)
json.dump({'seed': SEED, 'source': 'study_corpus/study_corpus.json', 'n': len(recs),
           'assignment': {p: sorted(cid(r) for r in halves[p]) for p in PEOPLE}},
          open(os.path.join(RDIR, 'retrieval_split.json'), 'w'), indent=1)
print(f"{len(recs)} included records -> {XLSX}")
for p in PEOPLE:
    nd = sum(1 for r in halves[p] if not (r.get('doi_raw') or r.get('doi')))
    print(f"  {p:8s} {len(halves[p]):4d} records  ({nd} without DOI)")
