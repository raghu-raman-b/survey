#!/usr/bin/env python3
"""Prepare the full-text screening corpus from the completed retrieval tracker.

Inputs:  study_corpus/study_corpus.json               the 880 records kept at title/abstract
         retrieval/retrieval_tracker_complete.xlsx    who retrieved what, status and notes
Output:  fulltext/fulltext_corpus.json                load this into tools/fulltext_screening.html
         prisma_stats.json  ('retrieval' block)       provisional retrieval counts for the diagram

Each screener screens the papers on their own sheet of the tracker. Some decisions are
pre-set from what is already known, so nobody has to click through them:
  - marked Not retrievable with a note like "Short Paper", "Abstract", "Demo", "WiP"
      -> exclude E1 (seen, not a full paper); hidden from the tool
  - marked Not retrievable for any other reason, or with no note
      -> not retrieved; hidden from the tool
  - acquired, but 4 pages or fewer by the page range
      -> exclude E1; listed under "Pre-set" in the tool, where it can be overridden

Usage:  python3 11_fulltext_corpus.py
"""
import json, os, re, datetime
from openpyxl import load_workbook

OUT  = '/home/prex_san/Documents/survey/final_corpus'
SRC  = os.path.join(OUT, 'study_corpus', 'study_corpus.json')
XLSX = os.path.join(OUT, 'retrieval', 'retrieval_tracker_complete.xlsx')
DST  = os.path.join(OUT, 'fulltext', 'fulltext_corpus.json')
STATS = os.path.join(OUT, 'prisma_stats.json')

recs = {f"C{r['cid']:04d}": r for r in json.load(open(SRC))}
wb = load_workbook(XLSX, data_only=True)
screeners = [s for s in wb.sheetnames if s != 'Summary']

# Notes that mean "we saw it, it is not a full paper" rather than "we could not get it".
FORMAT_NOTE = re.compile(r'short|abstract|demo|wip|work.in.progress|poster', re.I)

assignment, retrieval = {}, {}
for s in screeners:
    ids = []
    for row in wb[s].iter_rows(min_row=2, values_only=True):
        cid, status, note = row[0], row[6], (row[8] or '')
        if not cid or cid not in recs:
            continue
        ids.append(cid)
        note = str(note).strip()
        kind = ('acquired' if status == 'Acquired'
                else 'format' if FORMAT_NOTE.search(note)
                else 'no_access' if note else 'unexplained')
        retrieval[cid] = {'status': status or '', 'note': note, 'kind': kind}
    assignment[s] = ids

missing = set(recs) - set(retrieval)
if missing:
    raise SystemExit(f"{len(missing)} included records are not in the tracker, e.g. {sorted(missing)[:5]}")

def kws(r):
    k = r.get('keywords') or []
    if isinstance(k, str):                 # stored as a Python-list string in some records
        k = re.findall(r"'([^']+)'", k) or [k]
    return k[:12]

owner = {cid: s for s, ids in assignment.items() for cid in ids}

def n_pages(r):
    a, b = str(r.get('start_page') or ''), str(r.get('end_page') or '')
    return int(b) - int(a) + 1 if a.isdigit() and b.isdigit() and int(b) >= int(a) else None

def preset(cid, r):
    ret = retrieval[cid]
    if ret['kind'] == 'format':
        return {'decision': 'exclude', 'primary': 'E1', 'hidden': True,
                'reason': f"retrieval note: {ret['note']}"}
    if ret['kind'] in ('no_access', 'unexplained'):
        return {'decision': 'not_retrieved', 'primary': '', 'hidden': True,
                'reason': f"retrieval note: {ret['note']}" if ret['note'] else 'marked not retrievable, no note'}
    k = n_pages(r)
    if k is not None and k <= 4:
        return {'decision': 'exclude', 'primary': 'E1', 'hidden': False, 'reason': f"{k} pages"}
    return None

out = []
for cid in sorted(recs):
    r = recs[cid]
    out.append({
        'id': cid, 'type': r['type'], 'title': r['title'], 'authors': r.get('authors', []),
        'year': r['year'], 'venue': r.get('venue', ''), 'volume': r.get('volume', ''),
        'issue': r.get('issue', ''), 'start_page': r.get('start_page', ''),
        'end_page': r.get('end_page', ''), 'publisher': r.get('publisher', ''),
        'doi': r.get('doi_raw') or r.get('doi') or '', 'url': r.get('url', ''),
        'abstract': r.get('abstract', ''), 'keywords': kws(r),
        'found_in': r.get('found_in', []),
        'assigned_to': owner[cid], 'retrieval': retrieval[cid], 'preset': preset(cid, r),
    })

os.makedirs(os.path.dirname(DST), exist_ok=True)
json.dump({'tool': 'fulltext-desk', 'kind': 'fulltext-corpus', 'schema': 1,
           'created': datetime.date.today().isoformat(),
           'source': 'retrieval/retrieval_tracker_complete.xlsx',
           'screeners': assignment, 'records': out},
          open(DST, 'w'), indent=1, ensure_ascii=False)

# Provisional retrieval counts, used by 02_report.py until full-text screening is combined.
kinds = {}
for v in retrieval.values():
    kinds[v['kind']] = kinds.get(v['kind'], 0) + 1
S = json.load(open(STATS))
pre = [x['preset'] for x in out if x['preset']]
S['retrieval'] = {
    'sought': len(recs),
    'not_retrieved': sum(p['decision'] == 'not_retrieved' for p in pre),
    'preset_e1_retrieval_note': sum(p['decision'] == 'exclude' and p['hidden'] for p in pre),
    'preset_e1_short_pages': sum(p['decision'] == 'exclude' and not p['hidden'] for p in pre),
    'acquired': kinds.get('acquired', 0),
    'marked_not_retrievable': len(recs) - kinds.get('acquired', 0),
    'not_full_paper': kinds.get('format', 0),
    'no_access': kinds.get('no_access', 0),
    'unexplained': kinds.get('unexplained', 0),
    'updated': datetime.date.today().isoformat(),
}
json.dump(S, open(STATS, 'w'), indent=1)

print(f"{len(out)} records -> {DST}")
by_id = {x['id']: x for x in out}
for s, ids in assignment.items():
    hidden = sum(1 for c in ids if (by_id[c]['preset'] or {}).get('hidden'))
    shown_pre = sum(1 for c in ids if by_id[c]['preset'] and not by_id[c]['preset']['hidden'])
    print(f"  {s:8s} {len(ids)} assigned: {len(ids) - hidden} to screen ({shown_pre} of them pre-set E1 by page count), "
          f"{hidden} settled from retrieval notes")
print("retrieval:", {k: v for k, v in S['retrieval'].items() if k != 'updated'})
